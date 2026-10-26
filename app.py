"""Vinha: shared, email-authenticated visit management. WSGI application."""
import os, json, sqlite3, secrets, hashlib, time, re, io, zipfile, smtplib, ssl, calendar
from pathlib import Path
from datetime import date, datetime, timedelta
from email.message import EmailMessage
from urllib.parse import parse_qs
from http.cookies import SimpleCookie
from wsgiref.simple_server import make_server, WSGIRequestHandler
from zoneinfo import ZoneInfo
import openpyxl

ROOT = Path(__file__).parent
DATA = Path(os.getenv('DATA_DIR', str(ROOT / 'data')))
DEMO = os.getenv('APP_ENV') == 'demo'
BASE = os.getenv('BASE_URL', '').rstrip('/')
MAX_UPLOAD = 50 * 1024 * 1024
FREQUENCIES = {'Semanal': 7, 'Quinzenal': 14, 'Mensal': 0}

def connection():
    DATA.mkdir(parents=True, exist_ok=True)
    db = sqlite3.connect(DATA / 'vinha.sqlite3', timeout=20)
    db.row_factory = sqlite3.Row
    db.execute('PRAGMA foreign_keys=ON')
    db.execute('PRAGMA journal_mode=WAL')
    return db

def initialize():
    with connection() as db:
        db.executescript('''
        CREATE TABLE IF NOT EXISTS users(id INTEGER PRIMARY KEY,email TEXT UNIQUE NOT NULL,role TEXT NOT NULL CHECK(role IN ('admin','seller')),active INTEGER NOT NULL DEFAULT 1);
        CREATE TABLE IF NOT EXISTS sellers(id INTEGER PRIMARY KEY,name TEXT UNIQUE NOT NULL);
        CREATE TABLE IF NOT EXISTS grants(user_id INTEGER REFERENCES users(id),seller_id INTEGER REFERENCES sellers(id),PRIMARY KEY(user_id,seller_id));
        CREATE TABLE IF NOT EXISTS sessions(hash TEXT PRIMARY KEY,user_id INTEGER REFERENCES users(id),expires INTEGER,csrf TEXT);
        CREATE TABLE IF NOT EXISTS tokens(hash TEXT PRIMARY KEY,user_id INTEGER REFERENCES users(id),expires INTEGER);
        CREATE TABLE IF NOT EXISTS rates(key TEXT PRIMARY KEY,count INTEGER,until INTEGER);
        CREATE TABLE IF NOT EXISTS clients(id INTEGER PRIMARY KEY,cnpj TEXT NOT NULL,name TEXT NOT NULL,seller_id INTEGER NOT NULL REFERENCES sellers(id),channel TEXT NOT NULL,city TEXT,bairro TEXT,address TEXT,frequency TEXT NOT NULL,availability TEXT DEFAULT '',active INTEGER DEFAULT 1,central INTEGER DEFAULT 1,version INTEGER DEFAULT 1,UNIQUE(cnpj,seller_id));
        CREATE TABLE IF NOT EXISTS visits(id INTEGER PRIMARY KEY,client_id INTEGER NOT NULL REFERENCES clients(id),visit_date TEXT NOT NULL,status TEXT DEFAULT 'Agendada',version INTEGER DEFAULT 1,UNIQUE(client_id,visit_date));
        CREATE TABLE IF NOT EXISTS requests(id INTEGER PRIMARY KEY,client_id INTEGER REFERENCES clients(id),visit_id INTEGER REFERENCES visits(id),user_id INTEGER REFERENCES users(id),proposed_date TEXT NOT NULL,frequency TEXT NOT NULL,availability TEXT,reason TEXT,status TEXT DEFAULT 'Pendente',created TEXT,reviewed TEXT,admin_note TEXT,client_version INTEGER,visit_version INTEGER);
        CREATE TABLE IF NOT EXISTS audit(id INTEGER PRIMARY KEY,user_id INTEGER,action TEXT,payload TEXT,created TEXT);
        CREATE TABLE IF NOT EXISTS imports(id TEXT PRIMARY KEY,user_id INTEGER,expires INTEGER,payload TEXT);
        ''')
        columns={r[1] for r in db.execute('PRAGMA table_info(visits)')}
        for field,kind in [('checkin_at','TEXT'),('checkout_at','TEXT'),('checkin_user','INTEGER'),('checkin_lat','REAL'),('checkin_lon','REAL'),('checkout_lat','REAL'),('checkout_lon','REAL')]:
            if field not in columns:db.execute(f'ALTER TABLE visits ADD COLUMN {field} {kind}')
        admin = os.getenv('ADMIN_EMAIL', '').strip().lower()
        if admin:
            db.execute("INSERT INTO users(email,role) VALUES(?,'admin') ON CONFLICT(email) DO UPDATE SET role='admin',active=1", (email_address(admin),))
        if DEMO and not db.execute('SELECT 1 FROM clients LIMIT 1').fetchone():
            seed_demo(db)

def now(): return datetime.now(ZoneInfo('America/Sao_Paulo')).isoformat(timespec='seconds')
def today(): return datetime.now(ZoneInfo('America/Sao_Paulo')).date().isoformat()
def digest(value): return hashlib.sha256(value.encode()).hexdigest()
def dicts(rows): return [dict(r) for r in rows]
def audit(db, user, action, payload): db.execute('INSERT INTO audit(user_id,action,payload,created) VALUES(?,?,?,?)', (user['id'],action,json.dumps(payload,ensure_ascii=False),now()))

class Error(Exception):
    def __init__(self, status, message): self.status,self.message=status,message

def email_address(value):
    value=str(value).strip().lower()
    if len(value)>254 or not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',value): raise Error(400,'Informe um e-mail válido.')
    return value

def valid_date(value):
    try: return date.fromisoformat(str(value)).isoformat()
    except (ValueError,TypeError): raise Error(400,'Informe uma data válida.')

def frequency(value):
    if value not in FREQUENCIES: raise Error(400,'Escolha Semanal, Quinzenal ou Mensal.')
    return value

def rate(db, key, limit=5):
    ts=int(time.time()); row=db.execute('SELECT * FROM rates WHERE key=?',(key,)).fetchone()
    if row and row['until']>ts and row['count']>=limit: raise Error(429,'Muitas tentativas. Aguarde 15 minutos.')
    count=(row['count']+1) if row and row['until']>ts else 1
    until=row['until'] if row and row['until']>ts else ts+900
    db.execute('INSERT OR REPLACE INTO rates VALUES(?,?,?)',(key,count,until));db.commit()

def make_session(db, uid):
    token=secrets.token_urlsafe(32); csrf=secrets.token_urlsafe(24)
    db.execute('INSERT INTO sessions VALUES(?,?,?,?)',(digest(token),uid,int(time.time())+28800,csrf))
    return token

def send_access(db, user):
    if not BASE.startswith('https://') or not os.getenv('SMTP_HOST') or not os.getenv('MAIL_FROM'):
        raise Error(503,'O envio de acesso por e-mail ainda não foi configurado pelo administrador.')
    token=secrets.token_urlsafe(32)
    msg=EmailMessage();msg['Subject']='Seu acesso ao Vinha';msg['From']=os.environ['MAIL_FROM'];msg['To']=user['email']
    msg.set_content('Acesse sua agenda no Vinha. Este link expira em 15 minutos e só pode ser usado uma vez.\n\n'+BASE+'/?access='+token+'\n\nSe você não pediu este acesso, ignore a mensagem.')
    port=int(os.getenv('SMTP_PORT','587'))
    try:
        context=ssl.create_default_context()
        if port==465: server=smtplib.SMTP_SSL(os.environ['SMTP_HOST'],port,timeout=15,context=context)
        else:
            server=smtplib.SMTP(os.environ['SMTP_HOST'],port,timeout=15);server.starttls(context=context)
        with server:
            if os.getenv('SMTP_USER'): server.login(os.environ['SMTP_USER'],os.getenv('SMTP_PASSWORD',''))
            server.send_message(msg)
    except Exception: raise Error(503,'Não foi possível enviar o e-mail. Confira a configuração de envio.')
    db.execute('INSERT INTO tokens VALUES(?,?,?)',(digest(token),user['id'],int(time.time())+900))

def auth(db,environ):
    cookie=SimpleCookie()
    try: cookie.load(environ.get('HTTP_COOKIE',''))
    except Exception: return None
    raw=cookie.get('vinha_session')
    if not raw:return None
    row=db.execute('SELECT u.*,s.csrf FROM sessions s JOIN users u ON u.id=s.user_id WHERE s.hash=? AND s.expires>? AND u.active=1',(digest(raw.value),int(time.time()))).fetchone()
    return dict(row) if row else None

def admin(user):
    if user['role']!='admin':raise Error(403,'Acesso reservado ao administrador.')

def scope(user,alias='c'):
    return ('1=1',[]) if user['role']=='admin' else (f'{alias}.seller_id IN (SELECT seller_id FROM grants WHERE user_id=?)',[user['id']])

def client_access(db,user,cid):
    clause,args=scope(user)
    row=db.execute('SELECT c.*,s.name seller FROM clients c JOIN sellers s ON s.id=c.seller_id WHERE c.id=? AND '+clause,[cid]+args).fetchone()
    if not row:raise Error(404,'Cliente não encontrado.')
    return dict(row)

def visit_access(db,user,vid):
    v=db.execute('SELECT * FROM visits WHERE id=?',(vid,)).fetchone()
    if not v:raise Error(404,'Visita não encontrada.')
    c=client_access(db,user,v['client_id']);return dict(v),c

def ensure_capacity(db,c,day,exclude=None):
    if not c['active'] or not c['central']:raise Error(409,'Cliente inativo ou filial sem visita individual. Revise o cadastro.')
    if c['channel'] not in ['ON-TRADE','OFF-TRADE']:raise Error(409,'Defina o canal ON ou OFF antes de agendar.')
    units=1 if c['channel']=='ON-TRADE' else 2
    used=db.execute("SELECT COALESCE(SUM(CASE c.channel WHEN 'ON-TRADE' THEN 1 ELSE 2 END),0) FROM visits v JOIN clients c ON c.id=v.client_id WHERE c.seller_id=? AND v.visit_date=? AND v.status!='Cancelada' AND (? IS NULL OR v.id!=?)",(c['seller_id'],day,exclude,exclude)).fetchone()[0]
    if used+units>20:raise Error(409,'Limite diário excedido para este vendedor: 20 ON, 10 OFF ou carga mista proporcional. Escolha outra data.')

def recurring_days(start,freq,until):
    current=date.fromisoformat(start);end=date.fromisoformat(until);origin=current.day;out=[]
    if end<current or (end-current).days>93:raise Error(400,'O período deve terminar após o início e ter no máximo 93 dias.')
    while current<=end:
        out.append(current.isoformat())
        if freq=='Mensal':
            m=current.month%12+1;y=current.year+(current.month==12);current=date(y,m,min(origin,calendar.monthrange(y,m)[1]))
        else:current+=timedelta(days=FREQUENCIES[freq])
    return out

def is_business_day(day):
    d=date.fromisoformat(day)
    return d.weekday()<5 and (d.month,d.day) not in {(1,1),(4,21),(5,1),(9,7),(10,12),(11,2),(11,15),(11,20),(12,25)}

def norm(value):
    import unicodedata
    return re.sub(r'[^a-z0-9]','',unicodedata.normalize('NFKD',str(value or '')).encode('ascii','ignore').decode().lower())

def parse_xlsx(raw, requested_sheet=''):
    if len(raw)>MAX_UPLOAD:raise Error(413,'Limite de 50 MB por planilha.')
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if len(z.infolist())>2000 or sum(i.file_size for i in z.infolist())>512*1024*1024:raise Error(400,'Arquivo Excel excede os limites de leitura segura.')
            if any(i.file_size>30*1024*1024 for i in z.infolist() if 'sharedStrings' in i.filename):raise Error(400,'Planilha com excesso de textos compartilhados.')
        book=openpyxl.load_workbook(io.BytesIO(raw),read_only=True,data_only=True,keep_links=False)
    except Error:raise
    except Exception:raise Error(400,'Não consegui abrir o arquivo. Envie uma planilha .xlsx válida.')
    aliases={'cnpj':['cnpj','cnpjcpf','documento','cnpj14digitos'],'name':['nomedopn','cliente','clientecentral','razaosocial'], 'seller':['vendedor','representante'],'channel':['canal'],'city':['cidade','municipio','cidadedodestinatario','cidadedestinatario'],'bairro':['bairro'],'address':['endereco','enderecocompleto'],'number':['numero'],'date':['data','datadavisita','datavisita'],'frequency':['frequencia','frequenciasugerida'],'status':['situacaocadastral'],'person':['tipopessoa']}
    all_sheets=book.sheetnames
    candidates=[requested_sheet] if requested_sheet else all_sheets
    found=[]
    for name in candidates:
        if name not in all_sheets:raise Error(400,'Aba não encontrada.')
        sh=book[name];header=None;mapping=None
        for idx,row in enumerate(sh.iter_rows(max_row=40,max_col=40,values_only=True),1):
            keys=[norm(v) for v in row];m={k:next((i for i,v in enumerate(keys) if v in al),None) for k,al in aliases.items()}
            if m['cnpj'] is not None and m['name'] is not None:header=idx;mapping=m;break
        if header:found.append((name,header,mapping))
    if not found:raise Error(400,'Não localizei as colunas CNPJ e Cliente/Nome do PN.')
    if not requested_sheet:
        routes=[x for x in found if x[2]['date'] is not None]
        found=routes or [next((x for x in found if norm(x[0]) in ['planilha1','baseenviada','enderecoscnpj']),found[0])]
    records={};warnings=[];count=0
    for name,h,m in found:
        empty=0
        for rownum,row in enumerate(book[name].iter_rows(min_row=h+1,max_row=min((book[name].max_row or 20000),20000),max_col=40,values_only=True),h+1):
            def get(k):return row[m[k]] if m[k] is not None else None
            if not get('name'):
                empty+=1
                if empty>=200:break
                continue
            empty=0;count+=1
            if count>15000:raise Error(400,'Importe até 15 mil linhas por vez.')
            rawid=get('cnpj');digits=re.sub(r'\D','',str(int(rawid) if isinstance(rawid,(float,int)) else rawid or ''))
            if norm(get('person'))=='pessoafisica' or len(digits)>14 or not digits or (len(digits)<12 and norm(get('person'))!='pessoajuridica'):
                warnings.append(f'{name}, linha {rownum}: documento fora do escopo CNPJ.');continue
            cnpj=digits.zfill(14);seller=str(get('seller') or name).strip()
            channel=str(get('channel') or '').strip().upper().replace(' ','-')
            if channel not in ['ON-TRADE','OFF-TRADE']:channel='A CLASSIFICAR'
            freq=str(get('frequency') or ('Semanal' if channel=='ON-TRADE' else 'Mensal')).strip()
            if freq not in FREQUENCIES:freq='Mensal'
            key=cnpj+'|'+seller
            status=str(get('status') or '').upper()
            if key not in records:records[key]={'cnpj':cnpj,'name':str(get('name')).strip(),'seller':seller,'channel':channel,'city':str(get('city') or ''),'bairro':str(get('bairro') or ''),'address':str(get('address') or '')+(', '+str(get('number')) if get('number') else ''),'frequency':freq,'active':status not in ['BAIXADA','INAPTA','SUSPENSA'],'central':cnpj[8:12]=='0001','dates':[]}
            dt=get('date')
            if dt:
                try:
                    if isinstance(dt,datetime):day=dt.date().isoformat()
                    elif isinstance(dt,date):day=dt.isoformat()
                    elif isinstance(dt,(float,int)):day=openpyxl.utils.datetime.from_excel(dt).date().isoformat()
                    else:
                        try:day=date.fromisoformat(str(dt)).isoformat()
                        except ValueError:day=datetime.strptime(str(dt),'%d/%m/%Y').date().isoformat()
                    if day not in records[key]['dates']:records[key]['dates'].append(day)
                except Exception:warnings.append(f'{name}, linha {rownum}: data inválida, cliente importado sem essa visita.')
    book.close()
    return {'records':list(records.values()),'warnings':warnings,'sheets':[x[0] for x in found],'source_rows':count}

def commit_import(db,user,payload):
    imported=0;visits=0;warnings=list(payload['warnings']);duplicates=0
    for r in payload['records']:
        db.execute('INSERT OR IGNORE INTO sellers(name) VALUES(?)',(r['seller'],));sid=db.execute('SELECT id FROM sellers WHERE name=?',(r['seller'],)).fetchone()[0]
        old=db.execute('SELECT * FROM clients WHERE cnpj=? AND seller_id=?',(r['cnpj'],sid)).fetchone()
        if old:
            duplicates+=1;c=dict(old)
        else:
            cur=db.execute('INSERT INTO clients(cnpj,name,seller_id,channel,city,bairro,address,frequency,active,central) VALUES(?,?,?,?,?,?,?,?,?,?)',(r['cnpj'],r['name'],sid,r['channel'],r['city'],r['bairro'],r['address'],r['frequency'],int(r['active']),int(r['central'])))
            c=dict(db.execute('SELECT * FROM clients WHERE id=?',(cur.lastrowid,)).fetchone());imported+=1
        for day in r['dates']:
            if db.execute('SELECT 1 FROM visits WHERE client_id=? AND visit_date=?',(c['id'],day)).fetchone():continue
            try:
                if not is_business_day(day):raise Error(409,'fim de semana ou feriado nacional')
                ensure_capacity(db,c,day)
                db.execute('INSERT INTO visits(client_id,visit_date) VALUES(?,?)',(c['id'],day));visits+=1
            except Error as e:warnings.append(f"{r['name']} em {day}: visita não criada ({e.message}).")
    result={'clients_created':imported,'clients_preserved':duplicates,'visits_created':visits,'warnings':warnings}
    audit(db,user,'import_confirmed',result);return result

def dispatch(db,user,path,method,body,query,environ,raw):
    if path=='/api/config':return {'demo':DEMO,'mail_ready':bool(os.getenv('SMTP_HOST') and BASE.startswith('https://'))}
    if path=='/api/auth/request' and method=='POST':
        email=email_address(body.get('email',''));rate(db,'email:'+digest(email));rate(db,'ip:'+environ.get('REMOTE_ADDR',''),20)
        if not os.getenv('SMTP_HOST') or not BASE.startswith('https://'):raise Error(503,'Envio de acesso ainda não configurado. O administrador precisa configurar o serviço de e-mail.')
        u=db.execute('SELECT * FROM users WHERE email=? AND active=1',(email,)).fetchone()
        if u:send_access(db,u)
        return {'message':'Se este e-mail estiver autorizado, você receberá um link de acesso.'}
    if path=='/api/auth/verify' and method=='POST':
        rate(db,'verify:'+environ.get('REMOTE_ADDR',''),30)
        db.execute('BEGIN IMMEDIATE')
        row=db.execute('SELECT t.*,u.active FROM tokens t JOIN users u ON u.id=t.user_id WHERE t.hash=?',(digest(str(body.get('token',''))),)).fetchone()
        if not row or row['expires']<time.time() or not row['active']:raise Error(401,'Link inválido ou expirado. Solicite outro acesso.')
        db.execute('DELETE FROM tokens WHERE hash=?',(row['hash'],));return {'_session':make_session(db,row['user_id'])}
    if path=='/api/demo/login' and method=='POST' and DEMO:
        role='seller' if body.get('role')=='seller' else 'admin';u=db.execute('SELECT * FROM users WHERE role=? ORDER BY id LIMIT 1',(role,)).fetchone();return {'_session':make_session(db,u['id'])}
    if not user:raise Error(401,'Entre com seu e-mail autorizado.')
    if method!='GET' and not secrets.compare_digest(environ.get('HTTP_X_CSRF_TOKEN',''),user['csrf']):raise Error(403,'Sessão desatualizada. Recarregue a página.')
    if path=='/api/me':return {'user':{k:user[k] for k in ['id','email','role']},'csrf':user['csrf'],'today':today(),'demo':DEMO}
    if path=='/api/logout' and method=='POST':db.execute('DELETE FROM sessions WHERE user_id=?',(user['id'],));return {'_logout':True}
    if path=='/api/clients' and method=='GET':
        clause,args=scope(user)
        return {'clients':dicts(db.execute('SELECT c.*,s.name seller,(SELECT COUNT(*) FROM visits v WHERE v.client_id=c.id AND v.status!=\'Cancelada\') visits_count FROM clients c JOIN sellers s ON s.id=c.seller_id WHERE '+clause+' ORDER BY c.name',args))}
    if path=='/api/visits' and method=='GET':
        start=valid_date(query.get('from',[today()])[0]);end=valid_date(query.get('to',[start])[0]);clause,args=scope(user)
        rows=db.execute('SELECT v.*,c.name,c.cnpj,c.channel,c.city,c.bairro,c.address,c.frequency,c.availability,c.seller_id,s.name seller FROM visits v JOIN clients c ON c.id=v.client_id JOIN sellers s ON s.id=c.seller_id WHERE v.visit_date BETWEEN ? AND ? AND '+clause+' ORDER BY v.visit_date,s.name,c.city,c.bairro,c.name',[start,end]+args)
        return {'visits':dicts(rows)}
    if path=='/api/requests' and method=='GET':
        clause,args=scope(user)
        return {'requests':dicts(db.execute('SELECT r.*,c.name,c.seller_id,c.frequency current_frequency,s.name seller,v.visit_date current_date FROM requests r JOIN clients c ON c.id=r.client_id JOIN sellers s ON s.id=c.seller_id LEFT JOIN visits v ON v.id=r.visit_id WHERE '+clause+' ORDER BY r.id DESC',args))}
    if path=='/api/requests' and method=='POST':
        v,c=visit_access(db,user,int(body.get('visit_id',0)))
        if v['status']=='Cancelada':raise Error(409,'Visita cancelada.')
        day=valid_date(body.get('proposed_date'));freq=frequency(body.get('frequency'))
        if db.execute("SELECT 1 FROM requests WHERE visit_id=? AND status='Pendente'",(v['id'],)).fetchone():raise Error(409,'Já há uma solicitação pendente para esta visita.')
        cur=db.execute('INSERT INTO requests(client_id,visit_id,user_id,proposed_date,frequency,availability,reason,created,client_version,visit_version) VALUES(?,?,?,?,?,?,?,?,?,?)',(c['id'],v['id'],user['id'],day,freq,str(body.get('availability',''))[:1000],str(body.get('reason',''))[:1000],now(),c['version'],v['version']))
        audit(db,user,'request_created',{'id':cur.lastrowid});return {'id':cur.lastrowid,'message':'Solicitação enviada. A agenda só muda após a aprovação do administrador.'}
    if re.fullmatch(r'/api/visits/\d+/(checkin|checkout)',path) and method=='POST':
        vid=int(path.split('/')[3]);v,c=visit_access(db,user,vid);event=path.split('/')[-1]
        if v['visit_date']!=today():raise Error(409,'Check-in e check-out são permitidos para visitas de hoje. Solicite uma mudança de data se necessário.')
        if v['status'] in ['Cancelada','Realizada']:raise Error(409,'Esta visita já foi concluída ou cancelada.')
        lat=body.get('latitude');lon=body.get('longitude')
        if (lat is None)!=(lon is None):raise Error(400,'Localização incompleta.')
        if lat is not None and (not isinstance(lat,(int,float)) or not isinstance(lon,(int,float)) or not -90<=lat<=90 or not -180<=lon<=180):raise Error(400,'Localização inválida.')
        if event=='checkin':
            if v['checkin_at']:raise Error(409,'Check-in já registrado.')
            if db.execute("SELECT 1 FROM visits WHERE checkin_user=? AND checkin_at IS NOT NULL AND checkout_at IS NULL AND status='Em atendimento'",(user['id'],)).fetchone():raise Error(409,'Conclua o atendimento atual antes de iniciar o próximo.')
            db.execute("UPDATE visits SET status='Em atendimento',checkin_at=?,checkin_user=?,checkin_lat=?,checkin_lon=?,version=version+1 WHERE id=?",(now(),user['id'],lat,lon,vid))
        else:
            if not v['checkin_at']:raise Error(409,'Faça o check-in antes do check-out.')
            if v['checkin_user']!=user['id']:raise Error(403,'Somente quem iniciou o atendimento pode encerrá-lo.')
            db.execute("UPDATE visits SET status='Realizada',checkout_at=?,checkout_lat=?,checkout_lon=?,version=version+1 WHERE id=?",(now(),lat,lon,vid))
        audit(db,user,'visit_'+event,{'id':vid,'location_shared':lat is not None});return {'ok':True}
    if re.fullmatch(r'/api/clients/\d+',path) and method=='PATCH':
        admin(user);cid=int(path.split('/')[-1]);c=client_access(db,user,cid)
        if body.get('version')!=c['version']:raise Error(409,'Este cadastro foi alterado. Recarregue antes de salvar.')
        freq=frequency(body.get('frequency',c['frequency']));channel=body.get('channel',c['channel'])
        if channel not in ['ON-TRADE','OFF-TRADE','A CLASSIFICAR']:raise Error(400,'Canal inválido.')
        sid=int(body.get('seller_id',c['seller_id']))
        if not db.execute('SELECT 1 FROM sellers WHERE id=?',(sid,)).fetchone():raise Error(400,'Vendedor inválido.')
        if sid!=c['seller_id'] or channel!=c['channel']:raise Error(409,'Para mudar vendedor ou canal, cancele as visitas e ajuste o cadastro em uma importação revisada. Não alteramos a capacidade silenciosamente.')
        db.execute('UPDATE clients SET name=?,city=?,bairro=?,address=?,frequency=?,availability=?,active=?,central=?,version=version+1 WHERE id=?',(str(body.get('name',c['name']))[:200],str(body.get('city',c['city']))[:120],str(body.get('bairro',c['bairro']))[:120],str(body.get('address',c['address']))[:500],freq,str(body.get('availability',c['availability']))[:1000],int(bool(body.get('active',c['active']))),int(bool(body.get('central',c['central']))),cid))
        if not body.get('active',c['active']) or not body.get('central',c['central']):db.execute("UPDATE visits SET status='Cancelada',version=version+1 WHERE client_id=? AND status='Agendada' AND visit_date>=?",(cid,today()))
        audit(db,user,'client_updated',{'client_id':cid,'frequency':freq});return {'ok':True,'message':'Cadastro salvo. A frequência não altera automaticamente as datas existentes; use Gerar próximas visitas.'}
    if path=='/api/schedule' and method=='POST':
        admin(user);c=client_access(db,user,int(body.get('client_id',0)));start=valid_date(body.get('start'));freq=frequency(body.get('frequency',c['frequency']));until=valid_date(body.get('until',start));days=recurring_days(start,freq,until)
        skipped=[];created=[]
        for day in days:
            if not is_business_day(day):skipped.append({'date':day,'reason':'Fim de semana ou feriado nacional'});continue
            if db.execute('SELECT 1 FROM visits WHERE client_id=? AND visit_date=?',(c['id'],day)).fetchone():skipped.append({'date':day,'reason':'Já existe visita nesta data'});continue
            try:ensure_capacity(db,c,day)
            except Error as e:skipped.append({'date':day,'reason':e.message});continue
            db.execute('INSERT INTO visits(client_id,visit_date) VALUES(?,?)',(c['id'],day));created.append(day)
        if created:db.execute('UPDATE clients SET frequency=?,version=version+1 WHERE id=?',(freq,c['id']))
        audit(db,user,'schedule_created',{'client_id':c['id'],'created':created,'skipped':skipped});return {'created':created,'skipped':skipped}
    if re.fullmatch(r'/api/visits/\d+',path) and method=='PATCH':
        admin(user);v,c=visit_access(db,user,int(path.split('/')[-1]));day=valid_date(body.get('visit_date'))
        if body.get('version')!=v['version']:raise Error(409,'Esta visita mudou. Recarregue antes de editar.')
        if v['status'] in ['Em atendimento','Realizada']:raise Error(409,'Uma visita iniciada ou concluída não pode ser remarcada.')
        if not is_business_day(day):raise Error(409,'Escolha um dia útil. Feriados municipais precisam de conferência.')
        ensure_capacity(db,c,day,v['id']);freq=frequency(body.get('frequency',c['frequency']))
        db.execute('UPDATE visits SET visit_date=?,version=version+1 WHERE id=?',(day,v['id']));db.execute('UPDATE clients SET frequency=?,version=version+1 WHERE id=?',(freq,c['id']));audit(db,user,'visit_changed',{'id':v['id'],'date':day,'frequency':freq});return {'ok':True,'message':'Visita e frequência atualizadas. Outras datas já agendadas foram preservadas.'}
    if re.fullmatch(r'/api/requests/\d+/review',path) and method=='POST':
        admin(user);rid=int(path.split('/')[3]);r=db.execute('SELECT * FROM requests WHERE id=?',(rid,)).fetchone()
        if not r:raise Error(404,'Solicitação não encontrada.')
        if r['status']!='Pendente':raise Error(409,'Solicitação já analisada.')
        approved=body.get('decision')=='approve'
        if body.get('decision') not in ['approve','reject']:raise Error(400,'Decisão inválida.')
        if approved:
            v,c=visit_access(db,user,r['visit_id'])
            if v['status'] in ['Em atendimento','Realizada']:raise Error(409,'A visita já foi iniciada ou concluída. Não é possível remarcá-la por este pedido.')
            if v['version']!=r['visit_version'] or c['version']!=r['client_version']:raise Error(409,'A agenda ou o cadastro mudou desde o pedido. Edite a visita diretamente e rejeite este pedido antigo.')
            day=valid_date(body.get('date') or r['proposed_date']);freq=frequency(body.get('frequency') or r['frequency'])
            if not is_business_day(day):raise Error(409,'Escolha um dia útil para aprovar.')
            ensure_capacity(db,c,day,v['id'])
            db.execute('UPDATE visits SET visit_date=?,version=version+1 WHERE id=?',(day,v['id']));db.execute('UPDATE clients SET frequency=?,availability=?,version=version+1 WHERE id=?',(freq,r['availability'],c['id']))
        db.execute('UPDATE requests SET status=?,reviewed=?,admin_note=? WHERE id=?',('Aprovada' if approved else 'Recusada',now(),str(body.get('note',''))[:1000],rid));audit(db,user,'request_reviewed',{'id':rid,'approved':approved});return {'ok':True}
    if path=='/api/team' and method=='GET':
        admin(user);return {'sellers':dicts(db.execute('SELECT s.*,(SELECT COUNT(*) FROM clients c WHERE c.seller_id=s.id) client_count FROM sellers s ORDER BY name')),'users':dicts(db.execute('SELECT u.id,u.email,u.role,u.active,g.seller_id FROM users u LEFT JOIN grants g ON g.user_id=u.id ORDER BY u.email'))}
    if path=='/api/team' and method=='POST':
        admin(user);email=email_address(body.get('email',''));sid=int(body.get('seller_id',0))
        if not db.execute('SELECT 1 FROM sellers WHERE id=?',(sid,)).fetchone():raise Error(400,'Selecione um vendedor cadastrado.')
        existing=db.execute('SELECT * FROM users WHERE email=?',(email,)).fetchone()
        if existing and existing['role']=='admin':raise Error(409,'Este e-mail já é de administrador.')
        if existing and db.execute('SELECT 1 FROM grants WHERE user_id=? AND seller_id!=?',(existing['id'],sid)).fetchone():raise Error(409,'Este e-mail já está vinculado a outra carteira. Use um acesso por vendedor para manter a agenda individual.')
        db.execute("INSERT INTO users(email,role) VALUES(?,'seller') ON CONFLICT(email) DO UPDATE SET active=1",(email,));uid=db.execute('SELECT id FROM users WHERE email=?',(email,)).fetchone()[0]
        db.execute('INSERT OR IGNORE INTO grants VALUES(?,?)',(uid,sid));audit(db,user,'access_granted',{'user_id':uid,'seller_id':sid});return {'ok':True,'message':'E-mail autorizado. O vendedor pode solicitar seu link na tela de entrada.'}
    if re.fullmatch(r'/api/team/\d+',path) and method=='DELETE':
        admin(user);uid=int(path.split('/')[-1]);u=db.execute('SELECT * FROM users WHERE id=?',(uid,)).fetchone()
        if not u or u['role']=='admin':raise Error(400,'Este acesso não pode ser revogado por aqui.')
        db.execute('UPDATE users SET active=0 WHERE id=?',(uid,));db.execute('DELETE FROM sessions WHERE user_id=?',(uid,));db.execute('DELETE FROM tokens WHERE user_id=?',(uid,));audit(db,user,'access_revoked',{'user_id':uid});return {'ok':True}
    if path=='/api/import/preview' and method=='POST':
        admin(user);payload=parse_xlsx(raw,query.get('sheet',[''])[0]);key=secrets.token_urlsafe(24)
        db.execute('DELETE FROM imports WHERE expires<?',(int(time.time()),));db.execute('INSERT INTO imports VALUES(?,?,?,?)',(key,user['id'],int(time.time())+1800,json.dumps(payload,ensure_ascii=False)))
        return {'import_id':key,'clients':len(payload['records']),'visits':sum(len(r['dates']) for r in payload['records']),'sellers':sorted(set(r['seller'] for r in payload['records'])),'warnings':payload['warnings'],'sample':payload['records'][:8],'sheets':payload['sheets']}
    if path=='/api/import/confirm' and method=='POST':
        admin(user);row=db.execute('SELECT * FROM imports WHERE id=? AND user_id=? AND expires>?',(body.get('import_id'),user['id'],int(time.time()))).fetchone()
        if not row:raise Error(400,'Prévia expirada. Envie a planilha novamente.')
        result=commit_import(db,user,json.loads(row['payload']));db.execute('DELETE FROM imports WHERE id=?',(row['id'],));return result
    if path=='/api/audit' and method=='GET':
        admin(user);return {'audit':dicts(db.execute('SELECT a.*,u.email FROM audit a LEFT JOIN users u ON u.id=a.user_id ORDER BY a.id DESC LIMIT 100'))}
    raise Error(404,'Operação não encontrada.')

def application(environ,start_response):
    path=environ.get('PATH_INFO','/');method=environ.get('REQUEST_METHOD','GET');headers=[('X-Content-Type-Options','nosniff'),('Referrer-Policy','no-referrer'),('Cache-Control','no-store'),('Content-Security-Policy',"default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; connect-src 'self'; frame-ancestors 'self'; base-uri 'none'; form-action 'self'")]
    status=200
    try:
        if not path.startswith('/api/'):
            name={'/':'index.html','/app.js':'app.js','/style.css':'style.css'}.get(path)
            if not name:raise Error(404,'Página não encontrada.')
            content=(ROOT/'static'/name).read_bytes();ctype={'html':'text/html; charset=utf-8','js':'text/javascript; charset=utf-8','css':'text/css; charset=utf-8'}[name.split('.')[-1]]
        else:
            if method not in ['GET','POST','PATCH','DELETE']:raise Error(405,'Método não permitido.')
            if method!='GET':
                origin=environ.get('HTTP_ORIGIN','')
                expected=BASE or (environ.get('wsgi.url_scheme','http')+'://'+environ.get('HTTP_HOST',''))
                if origin and origin!=expected:raise Error(403,'Origem da requisição não autorizada.')
                if environ.get('HTTP_SEC_FETCH_SITE')=='cross-site':raise Error(403,'Origem não autorizada.')
            try:length=int(environ.get('CONTENT_LENGTH') or 0)
            except ValueError:raise Error(400,'Tamanho inválido.')
            if length>MAX_UPLOAD:raise Error(413,'Limite de 50 MB.')
            raw=environ['wsgi.input'].read(length) if length else b''
            body={}
            if raw and path!='/api/import/preview':
                if len(raw)>200000:raise Error(413,'Requisição muito grande.')
                try:body=json.loads(raw)
                except Exception:raise Error(400,'Dados inválidos.')
                if not isinstance(body,dict):raise Error(400,'Dados inválidos.')
            with connection() as db:
                user=auth(db,environ)
                if method!='GET':db.execute('BEGIN IMMEDIATE')
                result=dispatch(db,user,path,method,body,parse_qs(environ.get('QUERY_STRING','')),environ,raw)
                if '_session' in result:
                    token=result.pop('_session');secure='; Secure' if not DEMO else ''
                    headers.append(('Set-Cookie',f'vinha_session={token}; HttpOnly; SameSite=Lax; Path=/; Max-Age=28800{secure}'));result['ok']=True
                if result.pop('_logout',False):headers.append(('Set-Cookie','vinha_session=; HttpOnly; SameSite=Lax; Path=/; Max-Age=0'))
                content=json.dumps(result,ensure_ascii=False).encode();ctype='application/json; charset=utf-8'
    except Error as e:status=e.status;content=json.dumps({'error':e.message},ensure_ascii=False).encode();ctype='application/json; charset=utf-8'
    except sqlite3.IntegrityError:status=409;content=json.dumps({'error':'Já existe um registro com estes dados. Recarregue e confira a agenda.'}).encode();ctype='application/json'
    except Exception:status=500;content=b'{"error":"Falha interna. Nenhuma alteracao desta operacao foi confirmada."}';ctype='application/json'
    labels={200:'OK',400:'Bad Request',401:'Unauthorized',403:'Forbidden',404:'Not Found',405:'Method Not Allowed',409:'Conflict',413:'Payload Too Large',429:'Too Many Requests',500:'Internal Server Error',503:'Service Unavailable'}
    headers.extend([('Content-Type',ctype),('Content-Length',str(len(content)))])
    start_response(f'{status} {labels.get(status,"Error")}',headers);return [content]

def seed_demo(db):
    db.execute("INSERT OR IGNORE INTO users(id,email,role) VALUES(1,'gestao@exemplo.test','admin'),(2,'ana@exemplo.test','seller')")
    db.execute("INSERT OR IGNORE INTO sellers(id,name) VALUES(1,'Ana Martins'),(2,'Bruno Costa')")
    db.execute('INSERT OR IGNORE INTO grants VALUES(2,1)')
    samples=[('Bistrô do Jardim','ON-TRADE','Moinhos de Vento'),('Empório Aurora','OFF-TRADE','Bela Vista'),('Cantina da Praça','ON-TRADE','Bom Fim'),('Mercado do Vale','OFF-TRADE','Centro'),('Restaurante Oliva','ON-TRADE','Petrópolis'),('Casa do Vinho','OFF-TRADE','Auxiliadora')]
    for i,(name,channel,bairro) in enumerate(samples,1):
        sid=1 if i<=4 else 2;freq='Semanal' if channel=='ON-TRADE' else 'Quinzenal'
        db.execute('INSERT INTO clients(id,cnpj,name,seller_id,channel,city,bairro,address,frequency) VALUES(?,?,?,?,?,?,?,?,?)',(i,f'0000000{i}000100',name,sid,channel,'Porto Alegre',bairro,'Endereço fictício para demonstração',freq))
        db.execute('INSERT INTO visits(client_id,visit_date) VALUES(?,?)',(i,today()))
    db.execute('INSERT INTO requests(client_id,visit_id,user_id,proposed_date,frequency,availability,reason,created,client_version,visit_version) VALUES(2,2,2,?,?,?,?,?,1,1)',((date.fromisoformat(today())+timedelta(days=2)).isoformat(),'Mensal','Quartas, das 14h às 17h','Comprador atende somente à tarde.',now()))

initialize()
if __name__=='__main__':
    import argparse
    parser=argparse.ArgumentParser();parser.add_argument('--host',default='0.0.0.0');parser.add_argument('--port',type=int,default=int(os.getenv('PORT','3000')));args,_=parser.parse_known_args()
    class PrivateLogHandler(WSGIRequestHandler):
        def log_message(self,format,*args):pass
    make_server(args.host,args.port,application,handler_class=PrivateLogHandler).serve_forever()
